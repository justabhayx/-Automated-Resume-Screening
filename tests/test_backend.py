import pytest
from fastapi.testclient import TestClient
from docx import Document
from io import BytesIO
from PIL import Image

from backend.main import app
from backend.matcher import normalize_text, extract_skills_from_text, calculate_text_similarity, evaluate_candidate
from backend.parser import extract_text_from_docx, extract_text_from_image

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert "version" in data


def test_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert "message" in data


def test_text_normalization():
    raw = "  Python, FastAPI & Machine Learning! 100%  "
    norm = normalize_text(raw)
    assert "python" in norm
    assert "fastapi" in norm
    assert "machine learning" in norm
    assert "!" not in norm


def test_skill_extraction():
    text = "Experienced Senior Python Engineer with skills in FastAPI, Docker, PostgreSQL, and React. Strong leadership and Agile teamwork."
    skills = extract_skills_from_text(text)
    assert "python" in skills
    assert "fastapi" in skills
    assert "docker" in skills
    assert "postgresql" in skills
    assert "react" in skills
    assert "leadership" in skills
    assert "agile" in skills


def test_text_similarity():
    jd = "We are looking for a Python developer with FastAPI and PostgreSQL experience."
    resume_good = "Software engineer skilled in Python, FastAPI, and PostgreSQL databases."
    resume_bad = "Graphic designer proficient in Photoshop, Illustrator, and Figma."

    sim_good = calculate_text_similarity(jd, resume_good)
    sim_bad = calculate_text_similarity(jd, resume_bad)

    assert sim_good > sim_bad


def test_evaluate_candidate():
    jd = "Looking for Python, FastAPI, Docker, and PostgreSQL."
    resume = "Python developer with FastAPI, Docker, and PostgreSQL experience."

    eval_res = evaluate_candidate(jd, resume)
    assert "overall_score" in eval_res
    assert eval_res["overall_score"] > 50.0
    assert "python" in eval_res["matched_skills"]
    assert "explanation" in eval_res
    assert "disclaimer" in eval_res


def test_screen_endpoint_validation():
    # Missing job description
    response = client.post("/api/screen", data={})
    assert response.status_code == 422

    # Empty job description with dummy file
    doc = Document()
    doc.add_paragraph("Python developer resume.")
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    response = client.post(
        "/api/screen",
        data={"job_description": ""},
        files={"resumes": ("resume.docx", bio.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code in (400, 422)


def test_screen_endpoint_valid_docx():
    jd = "We need a Python developer with FastAPI and React experience."

    # Create sample docx 1 (Good match)
    doc1 = Document()
    doc1.add_paragraph("John Doe - Software Engineer")
    doc1.add_paragraph("Skills: Python, FastAPI, React, Docker, PostgreSQL.")
    bio1 = BytesIO()
    doc1.save(bio1)
    bio1.seek(0)

    # Create sample docx 2 (Partial/Low match)
    doc2 = Document()
    doc2.add_paragraph("Jane Smith - Marketing Specialist")
    doc2.add_paragraph("Skills: SEO, Content Writing, Photoshop.")
    bio2 = BytesIO()
    doc2.save(bio2)
    bio2.seek(0)

    response = client.post(
        "/api/screen",
        data={"job_description": jd},
        files=[
            ("resumes", ("john_resume.docx", bio1.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("resumes", ("jane_resume.docx", bio2.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total_submitted"] == 2
    assert data["total_successfully_screened"] == 2
    candidates = data["candidates"]
    assert len(candidates) == 2
    assert candidates[0]["filename"] == "john_resume.docx"
    assert candidates[0]["rank"] == 1
    assert candidates[1]["rank"] == 2
    assert candidates[0]["overall_score"] > candidates[1]["overall_score"]


def test_screen_endpoint_cors_preflight():
    """
    Verify that the Vercel frontend origin is allowed by CORS preflight.
    The browser sends an OPTIONS request with Origin header before the actual
    POST /api/screen call.
    """
    response = client.options(
        "/api/screen",
        headers={
            "Origin": "https://automated-resume-screening-six.vercel.app",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert response.status_code == 200
    # The CORS middleware should echo back the allowed origin
    assert "Access-Control-Allow-Origin" in response.headers


def test_screen_endpoint_unsupported_file_type():
    """
    Verify that unsupported file types are rejected and reported in the
    response's 'errors' array without crashing the whole request.
    """
    doc = Document()
    doc.add_paragraph("Valid resume content with skills: Python.")
    bio = BytesIO()
    doc.save(bio)

    response = client.post(
        "/api/screen",
        data={"job_description": "Looking for Python developer."},
        files=[
            ("resumes", ("unsupported.exe", b"MZ\x90\x00", "application/octet-stream")),
            ("resumes", ("valid.docx", bio.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total_errors"] >= 1
    assert any(
        e["filename"] == "unsupported.exe" and "Unsupported file extension" in e["error"]
        for e in data["errors"]
    )
    assert data["total_successfully_screened"] == 1
    assert data["candidates"][0]["filename"] == "valid.docx"


def test_screen_endpoint_response_contract():
    """
    Verify the exact JSON structure the frontend expects from /api/screen.
    This is the response contract that the frontend useScreen hook relies on.
    """
    doc = Document()
    doc.add_paragraph("Alex Johnson - Senior Backend Engineer")
    doc.add_paragraph("Skills: Python, FastAPI, Docker, PostgreSQL, AWS.")
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    response = client.post(
        "/api/screen",
        data={"job_description": "Senior Python Engineer with FastAPI, Docker, PostgreSQL experience."},
        files=[
            ("resumes", ("alex_resume.docx", bio.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
    )

    assert response.status_code == 200
    data = response.json()

    # Top-level fields
    assert "job_description_summary" in data
    assert "total_submitted" in data
    assert "total_successfully_screened" in data
    assert "total_errors" in data
    assert "candidates" in data
    assert "errors" in data

    assert data["total_submitted"] == 1
    assert data["total_successfully_screened"] == 1
    assert data["total_errors"] == 0

    candidate = data["candidates"][0]
    # Required fields the frontend CandidateAnalysis component reads
    assert "rank" in candidate
    assert "filename" in candidate
    assert "overall_score" in candidate
    assert "text_similarity" in candidate
    assert "skill_coverage" in candidate
    assert "matched_skills" in candidate
    assert "missing_skills" in candidate
    assert "extracted_skills" in candidate
    assert "explanation" in candidate
    assert "disclaimer" in candidate

    # Explanation sub-fields
    assert "strengths" in candidate["explanation"]
    assert "skill_gaps" in candidate["explanation"]
    assert "summary" in candidate["explanation"]

    # Scores are numbers in 0-100
    assert 0 <= candidate["overall_score"] <= 100
    assert 0 <= candidate["text_similarity"] <= 100
    assert 0 <= candidate["skill_coverage"] <= 100
